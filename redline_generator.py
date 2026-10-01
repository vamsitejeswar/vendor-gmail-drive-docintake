"""
Generates a redlined DOCX from a validation report.

Logic:
- Section 3 clauses with REVIEW NEEDED → find in doc → red highlight → yellow note below
- Section 4 missing clauses → yellow note at end of doc only (nothing to highlight)
"""

import io
import logging
import os
import re

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

logger = logging.getLogger(__name__)


# ── Parse Section 3 (clause-by-clause) ────────────────────────────────────────

def _parse_section3(report: str) -> list:
    """Return only REVIEW NEEDED clauses from Section 3."""
    items = []
    m = re.search(r"###\s*3\..*?Analysis(.*?)(?:###\s*4\.|$)", report, re.DOTALL | re.IGNORECASE)
    if not m:
        return items

    # Each clause block starts with "- **Clause:**"
    blocks = re.split(r"\n(?=\s*-\s*\*\*Clause:\*\*)", m.group(1))
    for blk in blocks:
        clause_m  = re.search(r"\*\*Clause:\*\*\s*(.+)", blk)
        verse_m   = re.search(r"\*\*Verse Standard:\*\*\s*(.+?)(?=\n\s*-\s*\*\*|\Z)", blk, re.DOTALL)
        thisdoc_m = re.search(r"\*\*This Document:\*\*\s*(.+?)(?=\n\s*-\s*\*\*|\Z)", blk, re.DOTALL)
        status_m  = re.search(r"\*\*Status:\*\*\s*(.+)", blk)

        if not clause_m or not status_m:
            continue
        status = status_m.group(1).strip()
        if "REVIEW NEEDED" not in status and "⚠️" not in status:
            continue

        items.append({
            "clause":    clause_m.group(1).strip(),
            "verse_std": re.sub(r"\s+", " ", verse_m.group(1)).strip() if verse_m else "",
            "this_doc":  re.sub(r"\s+", " ", thisdoc_m.group(1)).strip() if thisdoc_m else "",
            "status":    status,
        })
    return items


# ── Parse Section 4 (missing clauses) ─────────────────────────────────────────

def _parse_section4(report: str) -> list:
    """Return missing clauses — these go to the end of the doc."""
    items = []
    m = re.search(r"###\s*4\..*?Missing Clauses(.*?)(?:###\s*5\.|$)", report, re.DOTALL | re.IGNORECASE)
    if not m:
        return items
    blk = m.group(1)
    if "none" in blk[:80].lower():
        return items
    for line in blk.splitlines():
        cm = re.search(
            r"\*\*Clause:\*\*\s*(.+?)\s*—\s*\*\*Runbook Ref:\*\*\s*(\S+?)\s*—\s*\*\*Why it matters:\*\*\s*(.+)",
            line,
        )
        if cm:
            items.append({
                "clause":    cm.group(1).strip(),
                "verse_std": "",
                "this_doc":  "Clause not present in document.",
                "fix":       cm.group(3).strip(),
                "code":      cm.group(2).strip(),
            })
    return items


# ── Find matching paragraph ────────────────────────────────────────────────────

_CLAUSE_KEYWORDS = {
    "termination":               ["terminat"],
    "payment":                   ["payment", "commission", "net collection"],
    "confidentiality":           ["confidential information", "non-disclosure"],
    "intellectual property":     ["intellectual property", "work product"],
    "governing law":             ["governing law", "jurisdiction"],
    "indemnif":                  ["indemnif"],
    "anti-bribery":              ["anti-bribery", "prevention of corruption"],
    "anti-corruption":           ["anti-bribery", "prevention of corruption"],
    "publicity":                 ["trademark", "brand name", "publicity"],
    "data protection":           ["data protection", "personal data"],
    "privacy":                   ["data protection", "personal data"],
    "arbitration":               ["arbitration", "dispute resolution"],
    "dispute":                   ["arbitration", "dispute resolution"],
    "limitation of liability":   ["limitation of liability", "aggregate liability"],
    "survival":                  ["surviv"],
    "scope":                     ["scope of services", "scope of work"],
    "assignment":                ["assignment", "transfer of rights"],
    "non-solicitation":          ["non-solicitation", "solicit"],
    "effective date":            ["effective date", "this agreement is dated"],
    "parties":                   ["party of the first", "hereinafter referred"],
    "representations":           ["represent", "warrant"],
    "warranties":                ["warrant", "represent"],
    "signature":                 ["signed by", "authorized signatory", "in witness"],
}

_STOP_WORDS = {
    "clause", "clauses", "standard", "agreement", "contract", "section",
    "article", "parties", "party", "verse", "vendor", "document", "provision",
    "restriction", "requirement", "obligation", "rights", "issue", "include",
    "explicit", "specific", "missing", "noted", "right",
}

# If THIS DOC SAYS contains these → clause is absent, nothing to highlight
_ABSENT_SIGNALS = [
    "does not explicitly", "not explicitly", "no explicit", "no dedicated",
    "not present", "not included", "does not contain", "lacks", "absent",
    "no clause", "not found", "not addressed", "no provision", "not mentioned",
    "does not include", "no express", "not stated", "not defined",
    "does not restrict", "does not prohibit", "does not address",
    "no specific", "no separate", "no standalone",
]


def _is_absent_in_doc(this_doc: str) -> bool:
    """True when THIS DOC SAYS tells us the clause doesn't exist in the doc."""
    td = this_doc.lower()
    return any(sig in td for sig in _ABSENT_SIGNALS)


def _keywords_for_clause(clause_name: str) -> list:
    cl = clause_name.lower()
    for key, kws in _CLAUSE_KEYWORDS.items():
        if key in cl:
            return kws
    return [w.lower() for w in clause_name.split() if len(w) > 5 and w.lower() not in _STOP_WORDS]


def _extract_quotes(text: str) -> list:
    """Pull quoted phrases from THIS DOC text — most reliable paragraph anchor."""
    quotes = re.findall(r'["“”](.{15,}?)["“”]', text)
    quotes += re.findall(r"'(.{20,}?)'", text)
    return [q.strip() for q in quotes]


def _is_heading_para(para) -> bool:
    """True for titles/headings — skip these from highlighting."""
    txt = para.text.strip()
    if not txt or len(txt) < 5:
        return True
    if txt.isupper() and len(txt) < 100:
        return True
    if para.runs and all(r.bold for r in para.runs if r.text.strip()) and len(txt) < 100:
        return True
    return False


def _find_paragraph(paragraphs, clause_name: str, this_doc: str = "") -> int:
    # 1. Search by quoted text from THIS DOC — most accurate
    for quote in _extract_quotes(this_doc):
        snippet = quote.lower()[:50]
        for i, p in enumerate(paragraphs):
            txt = p.text.strip()
            if txt and len(txt) > 40 and not _is_heading_para(p) and snippet in txt.lower():
                return i

    # 2. Clause number from THIS DOC
    for num in re.findall(r'\b(\d+(?:\.\d+)+)\b', this_doc):
        for i, p in enumerate(paragraphs):
            txt = p.text.strip()
            if txt and not _is_heading_para(p) and num in txt:
                return i

    # 3. Whole-word keyword — only on substantive paragraphs (>60 chars)
    for kw in _keywords_for_clause(clause_name):
        pattern = re.compile(r'\b' + re.escape(kw) + r'\b', re.IGNORECASE)
        for i, p in enumerate(paragraphs):
            txt = p.text.strip()
            if txt and len(txt) > 60 and not _is_heading_para(p) and pattern.search(txt):
                return i

    return -1


# ── DOCX annotation helpers ────────────────────────────────────────────────────

def _redline_paragraph(para) -> None:
    """Red background on the conflicting clause paragraph."""
    for run in para.runs:
        rPr = run._element.get_or_add_rPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "FFCCCC")
        rPr.append(shd)
    pPr = para._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "FFCCCC")
    pPr.append(shd)


def _build_annotation(doc, item: dict, is_missing: bool = False) -> object:
    """Yellow note: Verse Standard vs This Document."""
    para = doc.add_paragraph()
    yellow = RGBColor(0xB8, 0x86, 0x00)

    label = "MISSING CLAUSE" if is_missing else "REVIEW NEEDED"
    code  = item.get("code", "")
    code_str = f" ({code})" if code else ""

    r1 = para.add_run(f"[{label}{code_str} — {item['clause']}]")
    r1.bold = True
    r1.font.color.rgb = yellow
    r1.font.size = Pt(9)

    if item.get("verse_std"):
        r2 = para.add_run(f"\n[Verse Standard: {item['verse_std']}]")
        r2.bold = True
        r2.font.color.rgb = yellow
        r2.font.size = Pt(8.5)

    if item.get("this_doc"):
        r3 = para.add_run(f"\n[This Doc: {item['this_doc']}]")
        r3.italic = True
        r3.font.color.rgb = yellow
        r3.font.size = Pt(8.5)

    if is_missing and item.get("fix"):
        r4 = para.add_run(f"\n[→ Fix: {item['fix']}]")
        r4.bold = True
        r4.font.color.rgb = yellow
        r4.font.size = Pt(8.5)

    # Yellow left border
    pPr = para._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), "B88600")
    pBdr.append(left)
    pPr.append(pBdr)

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "FFFBE6")
    pPr.append(shd)

    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "360")
    pPr.append(ind)

    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "120")
    pPr.append(spacing)

    elem = para._element
    doc.element.body.remove(elem)
    return elem


# ── Public API ─────────────────────────────────────────────────────────────────

def generate_redlined_docx(filename: str, file_bytes: bytes, validation_report: str) -> bytes | None:
    if os.path.splitext(filename)[1].lower() not in (".doc", ".docx"):
        logger.info("redline: skipping non-DOCX file")
        return None

    risky   = _parse_section3(validation_report)
    missing = _parse_section4(validation_report)

    if not risky and not missing:
        logger.info("redline: no items to annotate")
        return None

    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs  = doc.paragraphs
        used_idx    = set()
        footer_anno = []

        # Section 3 — highlight conflicting clauses inline
        for item in risky:
            anno = _build_annotation(doc, item, is_missing=False)

            # If THIS DOC SAYS the clause is absent/missing → footer only, nothing to highlight
            if _is_absent_in_doc(item.get("this_doc", "")):
                footer_anno.append(anno)
                logger.info(f"redline footer (absent): {item['clause']}")
                continue

            idx = _find_paragraph(paragraphs, item["clause"], item.get("this_doc", ""))
            if idx >= 0 and idx not in used_idx:
                used_idx.add(idx)
                _redline_paragraph(paragraphs[idx])
                paragraphs[idx]._element.addnext(anno)
                logger.info(f"redline inline: {item['clause']}")
            else:
                footer_anno.append(anno)
                logger.info(f"redline footer (no match): {item['clause']}")

        # Section 4 — missing clauses always at end
        for item in missing:
            footer_anno.append(_build_annotation(doc, item, is_missing=True))
            logger.info(f"redline footer (missing): {item['clause']}")

        if footer_anno:
            sep = doc.add_paragraph()
            sr = sep.add_run("─" * 20 + "  MISSING CLAUSES — VERSE LEGAL REVIEW  " + "─" * 20)
            sr.bold = True
            sr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            sr.font.size = Pt(8)
            for anno in footer_anno:
                doc.element.body.append(anno)

        out = io.BytesIO()
        doc.save(out)
        logger.info(f"Redlined DOCX: {len(used_idx)} inline, {len(footer_anno)} footer")
        return out.getvalue()

    except Exception as e:
        logger.error(f"generate_redlined_docx failed: {e}")
        return None
